#!/usr/bin/env python3
"""
Check a generated deck for layout faults, without rendering it.
===============================================================

    python presentation/check_layout.py [deck.pptx]

The deck is built programmatically from absolute coordinates, and there is no
PowerPoint or LibreOffice on this machine to render it. That is exactly the
situation where a slide silently ends up with text spilling out of its box or
running off the bottom of the slide, and nobody notices until it is on a
projector.

So we check the geometry directly:

    1. every shape lies inside the slide
    2. no text box overflows its own height, estimated from character count,
       font size and box width
    3. no two TEXT boxes overlap (a text box sitting on a background card is
       fine and expected; two text boxes on top of each other is not)

The text-height estimate is approximate - it assumes an average glyph width of
~0.5 em and wraps greedily - so it is tuned to flag genuine overflow rather
than to be exact. Treat a warning as "go and look at that slide".
"""

import sys
from collections import defaultdict

from pptx import Presentation
from pptx.util import Emu

SLIDE_W, SLIDE_H = 13.333, 7.5

# Average glyph advance as a fraction of font size, and line height as a
# multiple of it. Calibri sits near 0.48; we use a slightly generous value so
# the check errs toward reporting overflow rather than missing it.
GLYPH_W = 0.50
LINE_H = 1.22


def estimated_height(text, size_pt, width_in):
    """Roughly how tall this text will render, in inches, when wrapped."""
    if not text.strip():
        return 0.0
    chars_per_line = max(1, int(width_in * 72 / (size_pt * GLYPH_W)))
    lines = 0
    for paragraph in text.split("\n"):
        words, current = paragraph.split(), 0
        if not words:
            lines += 1
            continue
        line_len = 0
        for word in words:
            add = len(word) + (1 if line_len else 0)
            if line_len + add > chars_per_line and line_len:
                lines += 1
                line_len = len(word)
            else:
                line_len += add
        lines += 1
    return lines * size_pt * LINE_H / 72


def max_font(shape):
    """Largest run size in the shape, in points (default 18 if unset)."""
    sizes = [r.font.size.pt for p in shape.text_frame.paragraphs
             for r in p.runs if r.font.size]
    return max(sizes) if sizes else 18.0


def rects_overlap(a, b, tol=0.04):
    """Do two (l, t, w, h) rectangles overlap by more than `tol` inches?"""
    return (a[0] + a[2] - tol > b[0] and b[0] + b[2] - tol > a[0] and
            a[1] + a[3] - tol > b[1] and b[1] + b[3] - tol > a[1])


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else \
        "presentation/CGCNN_Elastic_Moduli.pptx"
    prs = Presentation(path)
    problems = defaultdict(list)

    for index, slide in enumerate(prs.slides, start=1):
        text_boxes = []

        for shape in slide.shapes:
            left, top = Emu(shape.left).inches, Emu(shape.top).inches
            width, height = Emu(shape.width).inches, Emu(shape.height).inches

            # --- 1. inside the slide? ---------------------------------
            if left < -0.02 or top < -0.02 or \
               left + width > SLIDE_W + 0.02 or top + height > SLIDE_H + 0.02:
                problems[index].append(
                    f"OFF-SLIDE  ({left:.2f},{top:.2f}) {width:.2f}x{height:.2f}"
                    f"  {getattr(shape, 'text', '')[:34]!r}")

            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            text = shape.text_frame.text

            # --- 2. does the text fit its box? ------------------------
            need = estimated_height(text, max_font(shape), width)
            if need > height + 0.28:            # 0.28in of slack
                problems[index].append(
                    f"OVERFLOW   needs ~{need:.2f}in, box {height:.2f}in"
                    f"  {text[:44]!r}")

            # Also: does the text, at its true height, run off the slide?
            if top + max(need, height) > SLIDE_H + 0.05:
                problems[index].append(
                    f"RUNS OFF   bottom at {top + need:.2f}in"
                    f"  {text[:44]!r}")

            text_boxes.append((left, top, width, max(height, need), text))

        # --- 3. text-on-text collisions -------------------------------
        for i in range(len(text_boxes)):
            for j in range(i + 1, len(text_boxes)):
                a, b = text_boxes[i], text_boxes[j]
                if rects_overlap(a[:4], b[:4]):
                    problems[index].append(
                        f"COLLISION  {a[4][:26]!r} <> {b[4][:26]!r}")

    total = sum(len(v) for v in problems.values())
    print(f"Checked {len(prs.slides)} slides in {path}")
    if not total:
        print("No layout problems found.")
        return 0

    print(f"{total} potential problem(s):\n")
    for index in sorted(problems):
        print(f"--- slide {index} ---")
        for message in problems[index]:
            print("   ", message)
    return 1


if __name__ == "__main__":
    sys.exit(main())
